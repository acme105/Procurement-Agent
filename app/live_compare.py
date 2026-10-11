"""Proposal comparison for Live RFQs: Phase 6 of rfx_copilot.ipynb as a background run with saved tables.

Proposals are the RFQ's active uploads (one file per organisation). A run extracts each proposal against the RFQ's own
lines and questions, then the math engine builds the comparison tables. Every run is saved under
STATE_DIR/live/rfqs/<RFQ>/analysis/<R-NNN>/ (manifest.json, extractions/*.json, *.parquet, derived/*.parquet) so the
page and the analyst read the same numbers.

Changes from Phase 6: RFQ read from its stored state (not re-extracted from the PDF); rows map to RFQ line numbers;
shortfall is per line; vendor name is the organisation from the upload; questionnaire style decides compliance
(none -> FAIL, blanket -> PASS with a flag, individual -> scored).
"""
from __future__ import annotations

import concurrent.futures
import io
import json
import re
import threading
from html import unescape
from pathlib import Path
from typing import List, Literal, Optional

import pandas as pd
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field

from . import config, live_rfx as L, llm
from .util import iso, sha256_file

router = APIRouter()
EXTRACT_PROMPT = "live_extract.v1"
PARAM_PROMPT = "live_param.v1"
TABLES = ["vendor_summary", "line_matrix", "cherry_pick", "line_items", "questionnaire", "rfq_lines", "parameters"]
DEFAULTS = {"fx_usd_inr": 84.0, "qa_threshold": 6.0}
_mlock = threading.Lock()


# ---- schemas (Phase 6, plus rfq_line, question_number and questionnaire_style) -----------------

class ExtractedLineItem(BaseModel):
    rfq_line: Optional[int] = Field(description="The RFQ line number (#) this row quotes. null if it matches no RFQ line.")
    item_description: str = Field(description="Description of the item as the vendor quoted it.")
    location: str = Field(description="Delivery location. Split into separate rows for each RFQ location if the vendor aggregates them.")
    quantity: float = Field(description="Quantity quoted for this specific location.")
    uom: str = Field(description="Unit of measure (e.g., 'Each', 'Box of 10')")
    unit_price: float = Field(description="Unit price quoted")
    currency: str = Field(description="Currency of the quote (e.g., 'INR', 'USD')")
    extraction_confidence: Literal["High", "Medium", "Low"] = Field(
        description="Low if reading a blurry photo, handwriting or unclear phrasing.")


class QuestionnaireAnswer(BaseModel):
    question_number: Optional[int] = Field(description="The RFQ question number this answers. null for one blanket answer covering all questions.")
    question: str
    answer: str
    confidence_score: float = Field(description="Rate the answer 1-10. Detailed, firm commitments score 10. Vague, evasive or conditional answers score lower.")
    confidence_reasoning: str = Field(description="A one sentence justification for the confidence score.")


class VendorProposal(BaseModel):
    vendor_name: str
    line_items: List[ExtractedLineItem]
    questionnaire_style: Literal["individual", "blanket", "none"] = Field(
        description="individual: answers the RFQ questions one by one. blanket: a single statement covering all questions "
                    "(e.g. 'We accept all RFQ terms'). none: does not respond to the questions at all.")
    questionnaire_answers: List[QuestionnaireAnswer]
    lead_time_days: Optional[int] = Field(description="Lead time in days, if mentioned.")
    payment_terms: Optional[str] = Field(description="e.g. Net 30, Net 15")
    freight_included: Optional[bool] = Field(description="True if freight/shipping is included in the price, False if extra.")
    bundle_conditions: Optional[str] = Field(description="Any conditional discounts (e.g., 'Discount applies only if awarded all items').")
    special_notes: Optional[str] = Field(description="Any other footnotes, exceptions, or RFQ terms rejected.")
    discount_percentage: float = Field(description="Discount percentage offered on the grand total. 0 if none.")


class ParamValue(BaseModel):
    value: Optional[str] = Field(description="The vendor's value for this parameter, short (e.g. '3 years', 'Yes', '12 days'). null if the proposal doesn't say.")
    evidence: Optional[str] = Field(description="A short verbatim quote from the proposal supporting the value. null if not stated.")
    confidence: Literal["High", "Medium", "Low"]


# ---- reading files (Phase 6's read_file_to_llm_parts, plus CSV, XLS and HTML) ---------------------

def _html_text(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d)>", "\n", raw)
    return re.sub(r"[ \t]+", " ", unescape(re.sub(r"<[^>]+>", " ", raw))).strip()


def read_parts(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext in (".txt", ".eml"):
        return [llm.text_part(path.read_text(errors="replace"))]
    if ext in (".html", ".htm"):
        return [llm.text_part(_html_text(path.read_text(errors="replace")))]
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
        return [llm.text_part(text)]
    if ext in (".xlsx", ".xls"):
        sheets = pd.read_excel(path, sheet_name=None)
        return [llm.text_part("\n\n".join(f"## {n}\n{df.to_markdown()}" for n, df in sheets.items()))]
    if ext == ".csv":
        return [llm.text_part(pd.read_csv(path).to_markdown())]
    if ext in (".jpeg", ".jpg", ".png"):
        return [llm.image_part(path.read_bytes(), mime="image/png" if ext == ".png" else "image/jpeg")]
    if ext == ".pdf":
        try:
            import pymupdf4llm
            md = pymupdf4llm.to_markdown(str(path))
            if len(md.strip()) > 100:
                return [llm.text_part(md)]
        except Exception:
            pass
        import pymupdf
        doc = pymupdf.open(str(path))
        return [llm.image_part(page.get_pixmap(dpi=150).tobytes("png"), mime="image/png") for page in doc]
    raise ValueError(f"{path.suffix} files can't be read for comparison")


# ---- proposals = active uploads -----------------------------------------------------------------

def proposals(st: dict) -> list[dict]:
    return [{"pid": f"P{u['n']:03d}", "n": u["n"], "org": u.get("org") or u.get("email") or "Unknown",
             "filename": u["filename"], "size": u["size"], "at": u["at"],
             "source": u.get("source") or ("email" if u.get("email") else "link")}
            for u in st["uploads"] if not u.get("superseded")]


def _upload_path(rid: str, st: dict, n: int) -> Path:
    u = next((u for u in st["uploads"] if u["n"] == n), None)
    if not u:
        raise HTTPException(404, "No such proposal")
    return L.rdir(rid) / "uploads" / u["path"]


def rfq_context(v: dict) -> str:
    lines = "\n".join(f"{l['n']} | {l['description']} | {l['location']} | {l['quantity']}" for l in v["lines"])
    qs = "\n".join(f"{i + 1}. {q}" for i, q in enumerate(v["questionnaire"])) or "(no questions)"
    return f"RFQ {v['id']}: {v['scope']}\n\nLines (# | item | location | quantity):\n{lines}\n\nQuestions for suppliers:\n{qs}"


def extract_instructions(v: dict) -> str:
    return f"""You are an expert procurement AI. Extract the vendor proposal details.

CRITICAL INSTRUCTIONS:
1. CONTEXT AWARENESS: Below is the official RFQ the buyer sent to the vendors:
{rfq_context(v)}

2. MAP LINES: For every quoted row set rfq_line to the RFQ line number it quotes (same item AND same location).
   If a vendor quotes one price for several locations, create one row per RFQ location, each with its own rfq_line.
   If a vendor explicitly restricts their service to specific locations, ONLY create rows for those locations.
   If a row matches no RFQ line, set rfq_line to null.

3. QUESTIONNAIRE: Set questionnaire_style. For 'individual', return one answer per question with its question_number.
   For 'blanket', return the single statement once with question_number null. For 'none', return no answers.
   Assign each answer a confidence_score (1-10).

4. CONFIDENCE: Set extraction_confidence to Low if reading from blurry handwriting or unclear phrasing.
"""


# ---- runs ---------------------------------------------------------------------------------------

def adir(rid: str) -> Path:
    return L.rdir(rid) / "analysis"


def run_dir(rid: str, run_id: str) -> Path:
    if not re.fullmatch(r"R-\d{3,}", run_id):
        raise HTTPException(404, "Unknown run")
    d = adir(rid) / run_id
    if not (d / "manifest.json").exists():
        raise HTTPException(404, "Unknown run")
    return d


def read_manifest(rid: str, run_id: str) -> dict:
    return json.loads((run_dir(rid, run_id) / "manifest.json").read_text())


def _write_manifest(rid: str, m: dict) -> None:
    d = adir(rid) / m["run_id"]
    d.mkdir(parents=True, exist_ok=True)
    (d / "manifest.json").write_text(json.dumps(m, ensure_ascii=False, indent=1))


def _update(rid: str, run_id: str, fn) -> dict:
    with _mlock:
        m = json.loads((adir(rid) / run_id / "manifest.json").read_text())
        fn(m)
        _write_manifest(rid, m)
        return m


def runs(rid: str) -> list[dict]:
    d = adir(rid)
    out = []
    for f in sorted(d.glob("R-*/manifest.json")) if d.exists() else []:
        m = json.loads(f.read_text())
        out.append({k: m.get(k) for k in ("run_id", "status", "created_at", "finished_at", "settings")} | {"proposals": len(m["proposals"])})
    return out


def latest(rid: str) -> Optional[str]:
    f = adir(rid) / "latest.txt"
    return f.read_text().strip() if f.exists() else None


def start_run(rid: str, settings: dict) -> dict:
    st = L.load(rid)
    if st["status"] != "FINAL":
        raise HTTPException(409, "Finalise the RFQ first")
    props = proposals(st)
    if not props:
        raise HTTPException(422, "No proposals yet. Add one, or wait for vendors to upload.")
    with _mlock:
        if any(r["status"] == "running" for r in runs(rid)):
            raise HTTPException(409, "A run is already in progress")
        n = max([int(r["run_id"][2:]) for r in runs(rid)], default=0) + 1
        prev = latest(rid)
        params = read_manifest(rid, prev).get("parameters", []) if prev else []
        m = {"run_id": f"R-{n:03d}", "rfq_id": rid, "status": "running", "created_at": iso(config.real_now()), "finished_at": None,
             "settings": {**DEFAULTS, **{k: float(v) for k, v in settings.items() if k in DEFAULTS and v is not None}},
             "model": config.MODEL, "prompt_version": EXTRACT_PROMPT, "parameters": params,
             "proposals": [{**p, "sha256": sha256_file(_upload_path(rid, st, p["n"])), "status": "queued", "error": None} for p in props]}
        _write_manifest(rid, m)
    threading.Thread(target=_execute, args=(rid, m["run_id"]), daemon=True).start()
    return m


def _extract_one(rid: str, run_id: str, p: dict, instructions: str, st: dict) -> None:
    def mark(status, error=None, **extra):
        def fn(m):
            for q in m["proposals"]:
                if q["pid"] == p["pid"]:
                    q.update(status=status, error=error, **extra)
        _update(rid, run_id, fn)
    mark("extracting")
    try:
        parts = read_parts(_upload_path(rid, st, p["n"]))
        parsed, run = llm.call("live_extract", EXTRACT_PROMPT, VendorProposal, instructions, parts,
                               doc_id=p["pid"], effort="medium", max_output_tokens=32000)
        out = run_dir(rid, run_id) / "extractions" / f"{p['pid']}.json"
        out.parent.mkdir(exist_ok=True)
        out.write_text(json.dumps({"pid": p["pid"], "org": p["org"], "filename": p["filename"],
                                   "proposal": parsed.model_dump(mode="json")}, ensure_ascii=False, indent=1))
        mark("done", cached=run.cached, llm_run=run.run_id)
    except Exception as e:  # unreadable file, model error: this proposal fails, the run continues
        mark("failed", f"{type(e).__name__}: {str(e)[:240]}")


def _execute(rid: str, run_id: str) -> None:
    st = L.load(rid)
    m = read_manifest(rid, run_id)
    instructions = extract_instructions(L.view(st))
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:
        list(ex.map(lambda p: _extract_one(rid, run_id, p, instructions, st), m["proposals"]))
    try:
        tables, summary = compute(rid, run_id, m["settings"])
        for name, df in tables.items():
            df.to_parquet(run_dir(rid, run_id) / f"{name}.parquet", index=False)
        for prm in m.get("parameters", []):          # parameters added in earlier runs are carried forward
            extract_parameter(rid, run_id, prm["name"], prm["description"], record=False)
        status = "done" if any(p["status"] == "done" for p in read_manifest(rid, run_id)["proposals"]) else "failed"
        _update(rid, run_id, lambda m: m.update(status=status, finished_at=iso(config.real_now()), summary=summary))
        if status == "done":
            (adir(rid) / "latest.txt").write_text(run_id)
    except Exception as e:
        _update(rid, run_id, lambda m: m.update(status="failed", finished_at=iso(config.real_now()), error=f"{type(e).__name__}: {str(e)[:300]}"))


def extractions(rid: str, run_id: str) -> list[dict]:
    d = run_dir(rid, run_id) / "extractions"
    return [json.loads(f.read_text()) for f in sorted(d.glob("P*.json"))] if d.exists() else []


# ---- math engine (Phase 6) ----------------------------------------------------------------------

def _ranges(nums: list[int]) -> str:
    """[69, 70, 71, 75] -> '69–71, 75'"""
    out, nums = [], sorted(nums)
    i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j > i else str(nums[i]))
        i = j + 1
    return ", ".join(out)


def _short(text: str, n: int = 70) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1] + "…"


def _uom_multiplier(uom: str) -> float:
    m = re.search(r"box of (\d+)", (uom or "").lower())
    return float(m.group(1)) if m else 1.0


def compute(rid: str, run_id: str, settings: dict, exclude: Optional[list[str]] = None) -> tuple[dict, dict]:
    """All comparison tables from a run's extractions. `exclude` drops vendors (organisation names) for scenarios."""
    v = L.view(L.load(rid))
    fx, qa_min = float(settings.get("fx_usd_inr", 84.0)), float(settings.get("qa_threshold", 6.0))
    rfq = pd.DataFrame([{"RFQ Line": l["n"], "Item": l["description"], "Location": l["location"],
                         "Qty": pd.to_numeric(l["quantity"], errors="coerce")} for l in v["lines"]])
    by_line = rfq.set_index("RFQ Line")
    req_locations = {str(x).strip().lower() for x in rfq["Location"]}
    excl = {L.org_key(x) for x in (exclude or [])}

    rows, q_rows, summ = [], [], []
    for e in extractions(rid, run_id):
        org, p = e["org"], e["proposal"]
        if L.org_key(org) in excl:
            continue
        for it in p["line_items"]:
            rate = fx if (it["currency"] or "").upper() == "USD" else 1.0
            mult = _uom_multiplier(it["uom"])
            norm_qty = it["quantity"] * mult
            per_piece = it["unit_price"] * rate / mult
            ln = it.get("rfq_line")
            matched = ln in by_line.index
            rows.append({"Vendor": org, "Proposal": e["pid"], "RFQ Line": ln if matched else None,
                         "RFQ Item": by_line.loc[ln, "Item"] if matched else None, "Vendor Item": it["item_description"],
                         "Location": it["location"], "Raw Qty": it["quantity"], "Raw UOM": it["uom"], "Raw Price": it["unit_price"],
                         "Currency": it["currency"], "Norm Qty (Pieces)": norm_qty, "Norm Unit Price (INR)": per_piece,
                         "Extended Cost (INR)": norm_qty * per_piece, "Extraction Confidence": it["extraction_confidence"],
                         "Matched": matched})
        answers = p["questionnaire_answers"]
        style = p["questionnaire_style"] if answers or p["questionnaire_style"] == "blanket" else "none"
        for qa in answers:
            q_rows.append({"Vendor": org, "Question #": qa.get("question_number"), "Question": qa["question"], "Answer": qa["answer"],
                           "Confidence Score": qa["confidence_score"], "Confidence Reasoning": qa["confidence_reasoning"], "Style": style})
        summ.append({"Vendor": org, "Proposal": e["pid"], "Q&A Style": style,
                     "Avg QA Confidence / 10": round(sum(a["confidence_score"] for a in answers) / len(answers), 1)
                     if answers and style == "individual" else None,
                     "Lead Time (Days)": p["lead_time_days"], "Payment Terms": p["payment_terms"],
                     "Freight Included": p["freight_included"], "Bundle Conditions": p["bundle_conditions"] or None,
                     "Special Notes": p["special_notes"] or None, "Discount %": p["discount_percentage"] or 0.0})

    cols = ["Vendor", "Proposal", "RFQ Line", "RFQ Item", "Vendor Item", "Location", "Raw Qty", "Raw UOM", "Raw Price", "Currency",
            "Norm Qty (Pieces)", "Norm Unit Price (INR)", "Extended Cost (INR)", "Extraction Confidence", "Matched"]
    li = pd.DataFrame(rows, columns=cols)
    li["RFQ Line"] = li["RFQ Line"].astype("Int64")
    qdf = pd.DataFrame(q_rows, columns=["Vendor", "Question #", "Question", "Answer", "Confidence Score", "Confidence Reasoning", "Style"])
    qdf["Question #"] = qdf["Question #"].astype("Int64")
    vs = pd.DataFrame(summ)
    m_li = li[li["Matched"]]

    # coverage and shortfall per RFQ line (not totals), then the Phase 6 knock-outs
    out = []
    for _, s in vs.iterrows():
        mine = m_li[m_li["Vendor"] == s["Vendor"]]
        quoted = mine.groupby("RFQ Line")["Norm Qty (Pieces)"].sum()
        need = rfq.set_index("RFQ Line")["Qty"]
        need_items = rfq.set_index("RFQ Line")["Item"]
        short = (need - quoted.reindex(need.index).fillna(0)).clip(lower=0)
        missing = [int(n) for n in need.index if n not in quoted.index]
        locs = {str(x).strip().lower() for x in mine["Location"]}
        ext = float(mine["Extended Cost (INR)"].sum())
        reasons, flags = [], []
        # which RFQ items are missing or under-quoted, named (not just counted)
        by_item: dict[str, list[int]] = {}
        for n in missing:
            by_item.setdefault(" ".join(str(need_items.loc[n]).split()), []).append(int(n))
        missing_text = "; ".join(f"{_short(item, 60)} (lines {_ranges(lines_)})" for item, lines_ in by_item.items())
        partial = [int(n) for n in short.index if short.loc[n] > 0 and n not in missing]
        if partial:
            missing_text += ("; " if missing_text else "") + f"quoted less than requested on lines {_ranges(partial)}"
        # rows that match no RFQ line: items the RFQ never asked for, vs requested items at extra locations
        all_rows = li[li["Vendor"] == s["Vendor"]]
        known = {str(x).lower() for x in all_rows[all_rows["Matched"]]["Vendor Item"]}
        off = all_rows[~all_rows["Matched"]]
        unrequested = off[~off["Vendor Item"].str.lower().isin(known)]
        extra_locs = off[off["Vendor Item"].str.lower().isin(known)]
        unrequested_text = "; ".join(f"{_short(item)} ({len(g)} rows)" for item, g in unrequested.groupby("Vendor Item", sort=False))
        if short.sum() > 0:
            reasons.append(f"Quantity shortfall: no quote for {missing_text}" if missing_text else "Quantity shortfall")
        if unrequested_text:
            reasons.append(f"Quoted items not in the RFQ: {unrequested_text}")
        if s["Q&A Style"] == "none":
            reasons.append("No questionnaire answers")
        elif s["Q&A Style"] == "individual" and pd.notna(s["Avg QA Confidence / 10"]) and s["Avg QA Confidence / 10"] < qa_min:
            reasons.append("Low QA score")
        if s["Q&A Style"] == "blanket":
            flags.append("Blanket Q&A answer")
        if pd.notna(s["Bundle Conditions"]) and str(s["Bundle Conditions"]).strip():
            flags.append("Bundle conditions")
        low = int((mine["Extraction Confidence"] == "Low").sum())
        if low:
            flags.append(f"{low} low-confidence rows")
        if len(extra_locs):
            where = sorted({str(x) for x in extra_locs["Location"]})
            flags.append(f"Quotes {len(extra_locs)} rows for locations the RFQ didn't list ({', '.join(where[:4])}{'…' if len(where) > 4 else ''})")
        out.append({**s.to_dict(), "Compliance Status": "FAIL" if reasons else "PASS", "Fail Reasons": "; ".join(reasons) or None,
                    "Flags": "; ".join(flags) or None, "Extended Cost (INR)": ext,
                    "Final Landed Cost (INR)": ext * (1 - (s["Discount %"] or 0) / 100),
                    "Lines Quoted": int(len(need) - len(missing)), "Lines in RFQ": int(len(need)),
                    "Missing Items": missing_text or None, "Items Not in RFQ": unrequested_text or None,
                    "Missing Lines": _ranges([int(n) for n in missing]) or None, "Quantity Shortfall": float(short.sum()),
                    "Delivers to All Locations": req_locations.issubset(locs)})
    order = ["Vendor", "Compliance Status", "Fail Reasons", "Flags", "Final Landed Cost (INR)", "Extended Cost (INR)", "Discount %",
             "Lines Quoted", "Lines in RFQ", "Missing Items", "Items Not in RFQ", "Missing Lines", "Quantity Shortfall", "Delivers to All Locations", "Q&A Style",
             "Avg QA Confidence / 10", "Lead Time (Days)", "Payment Terms", "Freight Included", "Bundle Conditions", "Special Notes", "Proposal"]
    vsum = pd.DataFrame(out, columns=order)
    if not vsum.empty:
        vsum = vsum.sort_values(["Compliance Status", "Final Landed Cost (INR)"], ascending=[False, True]).reset_index(drop=True)
        vsum.insert(0, "Rank", range(1, len(vsum) + 1))

    # line matrix: RFQ line x vendor, lowest normalised unit price
    best = m_li.groupby(["RFQ Line", "Vendor"])["Norm Unit Price (INR)"].min().unstack("Vendor") if not m_li.empty else pd.DataFrame()
    matrix = rfq.set_index("RFQ Line").join(best).reset_index()

    # cherry-pick: compliant vendors without bundle conditions (Phase 6), cheapest per RFQ line, RFQ quantities
    passing = vsum[vsum["Compliance Status"] == "PASS"] if not vsum.empty else vsum
    pool = passing[passing["Bundle Conditions"].isna()]["Vendor"].tolist() if not passing.empty else []
    basis = "PASS vendors without bundle conditions"
    if not pool:
        pool, basis = (passing["Vendor"].tolist(), "PASS vendors") if not passing.empty else (vsum["Vendor"].tolist() if not vsum.empty else [], "all vendors (none passed)")
    cp = []
    for _, r in rfq.iterrows():
        cands = (m_li[(m_li["RFQ Line"] == r["RFQ Line"]) & m_li["Vendor"].isin(pool)]
                 .sort_values("Norm Unit Price (INR)").drop_duplicates("Vendor"))
        top = cands.iloc[0] if len(cands) else None
        nxt = cands.iloc[1] if len(cands) > 1 else None
        cp.append({"RFQ Line": r["RFQ Line"], "Item": r["Item"], "Location": r["Location"], "Qty": r["Qty"],
                   "Vendor": top["Vendor"] if top is not None else None,
                   "Norm Unit Price (INR)": top["Norm Unit Price (INR)"] if top is not None else None,
                   "Cherry Pick Cost (INR)": top["Norm Unit Price (INR)"] * r["Qty"] if top is not None else None,
                   "Runner-up": nxt["Vendor"] if nxt is not None else None,
                   "Runner-up Price (INR)": nxt["Norm Unit Price (INR)"] if nxt is not None else None,
                   "Extraction Confidence": top["Extraction Confidence"] if top is not None else None})
    cpdf = pd.DataFrame(cp)
    total_cp = float(cpdf["Cherry Pick Cost (INR)"].sum()) if not cpdf.empty else 0.0
    best_single = float(passing["Final Landed Cost (INR)"].min()) if not passing.empty else None
    summary = {"cherry_pick_total": total_cp, "cherry_pick_basis": basis,
               "cherry_pick_uncovered_lines": int(cpdf["Vendor"].isna().sum()) if not cpdf.empty else 0,
               "best_single_source": best_single,
               "best_single_vendor": passing.iloc[0]["Vendor"] if not passing.empty else None,
               "savings_vs_best_single": (best_single - total_cp) if best_single is not None else None,
               "vendors": int(len(vsum)), "passing": int(len(passing)), "fx_usd_inr": fx, "qa_threshold": qa_min,
               "excluded": sorted(exclude or [])}
    params_f = run_dir(rid, run_id) / "parameters.parquet"
    params = pd.read_parquet(params_f) if params_f.exists() else pd.DataFrame(columns=["Parameter", "Description", "Vendor", "Value", "Evidence", "Confidence"])
    return {"vendor_summary": vsum, "line_matrix": matrix, "cherry_pick": cpdf, "line_items": li,
            "questionnaire": qdf, "rfq_lines": rfq, "parameters": params}, summary


# ---- custom parameters (added by the analyst) -----------------------------------------------------

def extract_parameter(rid: str, run_id: str, name: str, description: str, record: bool = True) -> pd.DataFrame:
    """Read one new criterion from every proposal in the run; save it to parameters.parquet."""
    st = L.load(rid)
    m = read_manifest(rid, run_id)
    instructions = (f"You are a procurement analyst. From this vendor proposal, extract the parameter '{name}': {description}\n"
                    "Answer only from the proposal. If it isn't stated, value and evidence are null.")

    def one(p):
        if p["status"] != "done":
            return None
        try:
            parsed, _ = llm.call("live_param", PARAM_PROMPT, ParamValue, instructions, read_parts(_upload_path(rid, st, p["n"])),
                                 doc_id=p["pid"], effort="low", max_output_tokens=2000)
            return {"Parameter": name, "Description": description, "Vendor": p["org"], "Value": parsed.value,
                    "Evidence": parsed.evidence, "Confidence": parsed.confidence}
        except Exception as e:
            return {"Parameter": name, "Description": description, "Vendor": p["org"], "Value": None,
                    "Evidence": f"Couldn't read: {str(e)[:120]}", "Confidence": "Low"}

    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:
        new = pd.DataFrame([r for r in ex.map(one, m["proposals"]) if r])
    f = run_dir(rid, run_id) / "parameters.parquet"
    with _mlock:
        old = pd.read_parquet(f) if f.exists() else pd.DataFrame(columns=new.columns)
        old = old[old["Parameter"].str.lower() != name.lower()] if not old.empty else old
        pd.concat([old, new], ignore_index=True).to_parquet(f, index=False)
    if record:
        _update(rid, run_id, lambda m: m.update(parameters=[x for x in m.get("parameters", []) if x["name"].lower() != name.lower()]
                                                + [{"name": name, "description": description}]))
    return new


# ---- tables -------------------------------------------------------------------------------------

def table_names(rid: str, run_id: str) -> list[str]:
    d = run_dir(rid, run_id)
    return [n for n in TABLES if (d / f"{n}.parquet").exists()] + [f"derived/{f.stem}" for f in sorted((d / "derived").glob("*.parquet"))]


def load_table(rid: str, run_id: str, name: str) -> pd.DataFrame:
    if not re.fullmatch(r"(derived/)?[a-z0-9_]{1,60}", name):
        raise HTTPException(404, f"Unknown table {name}")
    f = run_dir(rid, run_id) / f"{name}.parquet"
    if not f.exists():
        raise HTTPException(404, f"Unknown table {name}")
    return pd.read_parquet(f)


def save_derived(rid: str, run_id: str, name: str, df: pd.DataFrame) -> str:
    name = re.sub(r"[^a-z0-9_]+", "_", name.lower()).strip("_")[:60] or "table"
    d = run_dir(rid, run_id) / "derived"
    d.mkdir(exist_ok=True)
    df.to_parquet(d / f"{name}.parquet", index=False)
    return f"derived/{name}"


def records(df: pd.DataFrame, limit: int = 2000) -> list[dict]:
    return json.loads(df.head(limit).to_json(orient="records", date_format="iso"))


# ---- routes -------------------------------------------------------------------------------------

@router.get("/api/live/rfqs/{rid}/compare")
def compare_home(rid: str):
    st = L.load(rid)
    v = L.view(st)
    return {"id": rid, "scope": v["scope"], "status": v["status"], "lines": len(v["lines"]), "questions": v["questionnaire"],
            "proposals": proposals(st), "runs": runs(rid), "latest": latest(rid)}


@router.get("/api/live/compare/rfqs")
def compare_rfqs():
    """Finalised RFQs for the picker, newest first."""
    return [{"id": s["id"], "scope": s["rfx"]["scope"], "proposals": len(proposals(s))}
            for s in reversed(L.all_rfqs()) if s["status"] == "FINAL"]


@router.post("/api/live/rfqs/{rid}/proposals")
async def add_proposal(rid: str, org: str = Form(...), file: UploadFile = File(...)):
    rec = L.store_upload(rid, org, None, "buyer", file.filename, await file.read(L.MAX_UPLOAD + 1))
    return {"saved": rec["filename"], "replaced": rec["replaced"]}


@router.delete("/api/live/rfqs/{rid}/proposals/{n}")
def remove_proposal(rid: str, n: int):
    with L._lock:
        st = L.load(rid)
        u = next((u for u in st["uploads"] if u["n"] == n and not u.get("superseded")), None)
        if not u:
            raise HTTPException(404, "No such proposal")
        u["superseded"], u["removed_at"] = "removed", iso(config.real_now())
        L.save(st)
    return {"ok": True}


class RunIn(BaseModel):
    fx_usd_inr: Optional[float] = None
    qa_threshold: Optional[float] = None


@router.post("/api/live/rfqs/{rid}/runs")
def post_run(rid: str, body: RunIn):
    return start_run(rid, body.model_dump())


@router.get("/api/live/rfqs/{rid}/runs/{run_id}")
def get_run(rid: str, run_id: str):
    m = read_manifest(rid, run_id)
    out = {"manifest": m, "tables": {}}
    if m["status"] != "running":
        for name in table_names(rid, run_id):
            out["tables"][name] = records(load_table(rid, run_id, name))
    return out


@router.get("/api/live/rfqs/{rid}/runs/{run_id}/tables/{name:path}")
def get_table(rid: str, run_id: str, name: str, format: str = "json"):
    df = load_table(rid, run_id, name)
    fname = f"{rid}-{run_id}-{name.replace('/', '-')}"
    if format == "csv":
        return Response(df.to_csv(index=False), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})
    if format == "xlsx":
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})
    return records(df)


# ---- proposal viewer ------------------------------------------------------------------------------

@router.get("/api/live/rfqs/{rid}/proposals/{n}/view")
def view_proposal(rid: str, n: int):
    st = L.load(rid)
    path = _upload_path(rid, st, n)
    u = next(u for u in st["uploads"] if u["n"] == n)
    ext = path.suffix.lower()
    base = {"n": n, "org": u.get("org"), "filename": u["filename"], "file_url": f"/api/live/rfqs/{rid}/uploads/{n}"}
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(str(path)) as doc:
            return base | {"kind": "pages", "pages": [f"/api/live/rfqs/{rid}/proposals/{n}/pages/{i + 1}.png" for i in range(len(doc))]}
    if ext in (".png", ".jpg", ".jpeg"):
        return base | {"kind": "pages", "pages": [base["file_url"]]}
    if ext in (".xlsx", ".xls", ".csv"):
        df = pd.read_csv(path) if ext == ".csv" else pd.read_excel(path)
        return base | {"kind": "table", "columns": [str(c) for c in df.columns], "rows": json.loads(df.head(500).to_json(orient="values"))}
    try:
        part = read_parts(path)[0]
        return base | {"kind": "text", "text": part.get("text", "")[:200000]}
    except Exception as e:
        return base | {"kind": "text", "text": f"This file can't be previewed ({e}). Download it instead."}


@router.get("/api/live/rfqs/{rid}/proposals/{n}/pages/{p}.png")
def proposal_page(rid: str, n: int, p: int):
    st = L.load(rid)
    path = _upload_path(rid, st, n)
    out = L.rdir(rid) / "uploads" / ".pages" / f"{n:03d}-{p}.png"
    if not out.exists():
        import pymupdf
        with pymupdf.open(str(path)) as doc:
            if not 1 <= p <= len(doc):
                raise HTTPException(404, "No such page")
            out.parent.mkdir(exist_ok=True)
            doc[p - 1].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(str(out))
    return FileResponse(out, media_type="image/png")


# ---- the RFQ itself, in the same viewer ---------------------------------------------------------

@router.get("/api/live/rfqs/{rid}/rfq/view")
def view_rfq(rid: str):
    st = L.load(rid)
    f = L.rdir(rid) / "rfq.pdf"
    if st["status"] != "FINAL" or not f.exists():
        raise HTTPException(409, "Finalise the RFQ first")
    import pymupdf
    with pymupdf.open(str(f)) as doc:
        pages = [f"/api/live/rfqs/{rid}/rfq/pages/{i + 1}.png" for i in range(len(doc))]
    return {"n": 0, "org": f"{rid} · the RFQ sent to vendors", "filename": f"{rid}.pdf", "file_url": f"/api/live/rfqs/{rid}/pdf",
            "kind": "pages", "pages": pages}


@router.get("/api/live/rfqs/{rid}/rfq/pages/{p}.png")
def rfq_page(rid: str, p: int):
    f = L.rdir(rid) / "rfq.pdf"
    if not f.exists():
        raise HTTPException(404, "No RFQ PDF")
    out = L.rdir(rid) / ".rfq-pages" / f"{p}-{int(f.stat().st_mtime)}.png"     # keyed by file time: reprints refresh it
    if not out.exists():
        import pymupdf
        with pymupdf.open(str(f)) as doc:
            if not 1 <= p <= len(doc):
                raise HTTPException(404, "No such page")
            out.parent.mkdir(exist_ok=True)
            doc[p - 1].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(str(out))
    return FileResponse(out, media_type="image/png")
